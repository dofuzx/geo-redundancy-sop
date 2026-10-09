#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown -> docx converter (headings, bold, superscript, lists, tables,
images, blockquotes; skips --- rules)."""
import re
import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "Manuscript_BIB_CRC_GEO_redundancy.md")
OUT = os.path.join(BASE, "Manuscript_BIB_CRC_GEO_redundancy.docx")
IMG_RE = re.compile(r"^!\[[^\]]*\]\(([^)]+)\)\s*$")
BQ_RE = re.compile(r"^>\s?(.*)$")


def add_runs(p, text):
    for seg in re.split(r'(\*\*.+?\*\*)', text):
        if seg.startswith("**") and seg.endswith("**") and len(seg) >= 4:
            bold = seg[2:-2]
            for s in re.split(r'(\^.+?\^)', bold):
                if s.startswith("^") and s.endswith("^") and len(s) >= 2:
                    r = p.add_run(s[1:-1]); r.bold = True; r.font.superscript = True
                elif s:
                    r = p.add_run(s); r.bold = True
        else:
            for s in re.split(r'(\^.+?\^)', seg):
                if s.startswith("^") and s.endswith("^") and len(s) >= 2:
                    r = p.add_run(s[1:-1]); r.font.superscript = True
                elif s:
                    r = p.add_run(s)


def is_table_row(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def parse_table(rows):
    data = []
    for row in rows:
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        data.append(cells)
    return data[0], data[2:]


def main():
    doc = Document()
    base = doc.styles["Normal"]
    base.font.name = "Times New Roman"
    base.font.size = Pt(11)

    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        if line.startswith("# "):
            p = doc.add_heading(level=0)
            add_runs(p, line[2:])
            i += 1
        elif line.startswith("## "):
            p = doc.add_heading(level=1)
            add_runs(p, line[3:])
            i += 1
        elif line.startswith("### "):
            p = doc.add_heading(level=2)
            add_runs(p, line[4:])
            i += 1
        elif line.strip() == "---":
            i += 1
        elif IMG_RE.match(line.strip()):
            rel = IMG_RE.match(line.strip()).group(1)
            path = os.path.join(BASE, rel)
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Inches(6.0))
            else:
                p = doc.add_paragraph()
                add_runs(p, "[missing figure: %s]" % rel)
            i += 1
        elif BQ_RE.match(line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            r = p.add_run(BQ_RE.match(line).group(1))
            r.italic = True
            i += 1
        elif is_table_row(line):
            block = []
            while i < n and is_table_row(lines[i]):
                block.append(lines[i]); i += 1
            header, rows = parse_table(block)
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for j, h in enumerate(header):
                add_runs(t.rows[0].cells[j].paragraphs[0], h)
            for rrow in rows:
                cells = t.add_row().cells
                for j, c in enumerate(rrow):
                    add_runs(cells[j].paragraphs[0], c)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
            i += 1
        elif line.strip() == "":
            i += 1
        else:
            p = doc.add_paragraph()
            add_runs(p, line)
            i += 1

    doc.save(OUT)
    print("Saved:", OUT, os.path.getsize(OUT), "bytes")


if __name__ == "__main__":
    main()
